import os
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

async def run_ppt_to_pdf(file_paths: list[str], temp_dir: str, orig_filename: str, unique_suffix: str) -> str:
    input_path = file_paths[0]
    output_filename = f"{orig_filename}_{unique_suffix}.pdf"
    output_path = os.path.join(temp_dir, output_filename)
    
    prs = Presentation(input_path)
    pdf_doc = SimpleDocTemplate(output_path, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []
    
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.text.strip():
                        story.append(Paragraph(paragraph.text, styles['Normal']))
                        story.append(Spacer(1, 4))
        story.append(Spacer(1, 12))
        
    pdf_doc.build(story)
    return output_path
