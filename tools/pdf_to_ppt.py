import os
import fitz
import uuid
from pptx import Presentation
from pptx.util import Inches

async def run_pdf_to_ppt(file_paths: list[str], temp_dir: str, orig_filename: str, unique_suffix: str) -> str:
    input_path = file_paths[0]
    output_filename = f"{orig_filename}_{unique_suffix}.pptx"
    output_path = os.path.join(temp_dir, output_filename)
    
    prs = Presentation()
    doc = fitz.open(input_path)
    
    for page in doc:
        pix = page.get_pixmap()
        img_path = os.path.join(temp_dir, f"slide_{uuid.uuid4().hex[:6]}.png")
        pix.save(img_path)
        
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(img_path, 0, 0, width=Inches(10))
        
        if os.path.exists(img_path):
            os.remove(img_path)
            
    prs.save(output_path)
    doc.close()
    return output_path
